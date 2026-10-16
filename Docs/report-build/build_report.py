"""Fills a copy of the course template with content.py -> Docs/AWS_Cloud_Project_Report_Final.docx.

    python build_report.py

The template's page setup, styles, headings, tables and footer are kept: placeholders are located by
their exact template text (the build fails loudly if one is missing), replaced in place, and the
template's instructions are removed. Word (update_fields.ps1) then refreshes the table of contents
and exports the PDF.
"""
import copy
import re
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from PIL import Image

import content as C

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
TEMPLATE = ROOT / 'Docs' / 'AWS_Cloud_Project_Report_Template.docx'
OUT = ROOT / 'Docs' / 'AWS_Cloud_Project_Report_Final.docx'

doc = docx.Document(TEMPLATE)
BODY = doc.element.body
FIGS = {k: (path, cap) for k, path, cap in C.FIGURES}
fig_numbers: dict[str, int] = {}


# ------------------------------------------------------------------------------------------------
# figure numbering follows the order in which figures are placed (computed before anything renders)
def collect_figs(blocks):
    for b in blocks:
        if b[0] == 'fig' and b[1] not in fig_numbers:
            fig_numbers[b[1]] = len(fig_numbers) + 1


def fx(text: str) -> str:
    return re.sub(r'\{F:([a-z0-9_]+)\}', lambda m: str(fig_numbers[m.group(1)]), text)


# ------------------------------------------------------------------------------------------------
# paragraph helpers
def paragraphs():
    return [Paragraph(p, doc._body) for p in BODY.iterchildren(qn('w:p'))]


def find(prefix: str) -> Paragraph:
    hits = [p for p in paragraphs() if p.text.strip().startswith(prefix)]
    if len(hits) != 1:
        sys.exit(f'template anchor not found exactly once ({len(hits)}): {prefix!r}')
    return hits[0]


def find_heading(text: str) -> Paragraph:
    hits = [p for p in paragraphs() if p.style.name.startswith('Heading') and p.text.strip() == text]
    if len(hits) != 1:
        sys.exit(f'heading not found exactly once ({len(hits)}): {text!r}')
    return hits[0]


def delete(p: Paragraph):
    p._p.getparent().remove(p._p)


def new_after(elm, style='Normal') -> Paragraph:
    p = OxmlElement('w:p')
    elm.addnext(p)
    para = Paragraph(p, doc._body)
    para.style = doc.styles[style]
    return para


def add_text(para: Paragraph, text: str, bold_lead=False):
    text = fx(text)
    if bold_lead:
        m = re.match(r'^(.{2,80}?[.:])(\s.*)$', text, re.S)
        if m:
            r = para.add_run(m.group(1))
            r.bold = True
            para.add_run(m.group(2))
            return
    para.add_run(text)


# Console screenshots are cropped to the area that carries the evidence (no side menus), so their
# text prints larger. Boxes are (left, top, right, bottom) in source pixels.
CROPS = {
    'cfn': (255, 55, 1600, 735), 'vpcmap': (455, 90, 2090, 1010), 'endpoints': (245, 95, 1600, 390),
    'eks': (285, 95, 1600, 730), 'ec2': (245, 95, 1600, 345), 'msk': (285, 95, 1600, 620),
    'neptune': (285, 95, 1600, 700), 'transformer': (0, 95, 1310, 600), 'scorer': (0, 95, 1310, 600),
    'sagemaker': (285, 95, 1600, 440), 'dynamodb': (285, 175, 1600, 405), 'sns': (285, 175, 1600, 770),
    'apigw': (280, 95, 1240, 600), 'cloudfront': (285, 95, 1600, 860), 'github': (180, 180, 1420, 960),
    'budget': (285, 195, 1600, 360), 'kms': (605, 180, 1600, 880), 'scorer_metrics': (20, 680, 1290, 1770),
}
LANDSCAPE = {'aws', 'system'}
CROP_DIR = HERE / 'figures' / 'crops'


def add_figure(after_elm, key):
    path, caption = FIGS[key]
    img = ROOT / path
    if not img.exists():
        sys.exit(f'missing figure image: {img}')
    if key in CROPS:
        CROP_DIR.mkdir(exist_ok=True)
        cropped = CROP_DIR / f'{key}.png'
        Image.open(img).crop(CROPS[key]).save(cropped)
        img = cropped
    w, h = Image.open(img).size
    max_w, max_h = (8.6, 5.35) if key in LANDSCAPE else (6.5, 7.6)
    width = min(max_w, max_h * w / h)
    pic = new_after(after_elm)
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.keep_with_next = True
    pic.paragraph_format.space_before = Pt(6)
    pic.add_run().add_picture(str(img), width=Inches(width))
    cap = new_after(pic._p)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    r = cap.add_run(f'Figure {fig_numbers[key]}. ')
    r.bold = True
    r2 = cap.add_run(fx(caption))
    for run in (r, r2):
        run.font.size = Pt(11)
    return cap._p


def render(blocks, after_elm):
    """Inserts blocks after after_elm; returns the last inserted element."""
    last = after_elm
    for b in blocks:
        kind = b[0]
        if kind in ('p', 'pl'):
            para = new_after(last)
            add_text(para, b[1], bold_lead=(kind == 'pl'))
            last = para._p
        elif kind == 'b':
            for item in b[1]:
                para = new_after(last, 'List Bullet')
                add_text(para, item)
                last = para._p
        elif kind == 'steps':
            for i, item in enumerate(b[1], 1):
                para = new_after(last)
                para.paragraph_format.left_indent = Inches(0.3)
                para.paragraph_format.first_line_indent = Inches(-0.3)
                r = para.add_run(f'{i}.\t')
                r.bold = True
                add_text(para, item)
                last = para._p
        elif kind == 'fig':
            last = add_figure(last, b[1])
        else:
            sys.exit(f'unknown block {kind}')
    return last


def replace(prefix: str, blocks):
    anchor = find(prefix)
    render(blocks, anchor._p)
    delete(anchor)


def set_text(p: Paragraph, text: str, bold_lead=False):
    """Replaces a paragraph's text, keeping its style; first run formatting kept, colour/italic dropped."""
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    if p.runs:
        r0 = p.runs[0]
        r0.text = ''
        r0.italic = None
        r0.font.color.rgb = None
    if bold_lead:
        if p.runs:
            p.runs[0]._r.getparent().remove(p.runs[0]._r)
        add_text(p, text, bold_lead=True)
    elif p.runs:
        p.runs[0].text = fx(text)
    else:
        p.add_run(fx(text))


def page_break_before(p: Paragraph):
    br = OxmlElement('w:p')
    p._p.addprevious(br)
    Paragraph(br, doc._body).add_run().add_break(WD_BREAK.PAGE)


# ------------------------------------------------------------------------------------------------
# table helpers
def set_cell(cell, text: str, bold=False):
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    para = cell.paragraphs[0]
    for r in list(para.runs):
        r._r.getparent().remove(r._r)
    run = para.add_run(fx(text))
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    run.bold = True if bold else None


def fill_table(t, rows, header=None):
    tmpl_row = t.rows[1]._tr
    # remove all body rows, then add one clone of the first body row per data row
    for row in list(t.rows)[1:]:
        row._tr.getparent().remove(row._tr)
    for data in rows:
        tr = copy.deepcopy(tmpl_row)
        t._tbl.append(tr)
        row = t.rows[-1]
        if len(data) != len(row.cells):
            sys.exit(f'row width {len(data)} != table width {len(row.cells)}: {data[0]}')
        for cell, text in zip(row.cells, data):
            set_cell(cell, text)
    if header:
        for cell, text in zip(t.rows[0].cells, header):
            set_cell(cell, text, bold=True)
    # repeat the header row on each page; never split a row across pages
    for i, row in enumerate(t.rows):
        trPr = row._tr.get_or_add_trPr()
        if i == 0:
            el = OxmlElement('w:tblHeader')
            el.set(qn('w:val'), 'true')
            trPr.append(el)
        el = OxmlElement('w:cantSplit')
        el.set(qn('w:val'), 'true')
        trPr.append(el)


# ------------------------------------------------------------------------------------------------
# 0. figure order (document order of every figure block)
ORDERED_BLOCKS = [
    C.SYS_AWS, [('fig', 'aws')], [('fig', 'system')], C.SYS_COMPLETE, C.AWS_COST, C.METHOD, [('fig', 'gnn')],
    C.SCREENSHOTS, C.REPO_TEXT, C.RESULTS,
]
for blocks in ([('fig', 'aws')], [('fig', 'system')], C.AWS_COST, [('fig', 'gnn')], C.SCREENSHOTS, C.REPO_TEXT, C.RESULTS):
    collect_figs(blocks)
missing = [k for k in FIGS if k not in fig_numbers]
if missing:
    sys.exit(f'figures defined but never placed: {missing}')

# 1. cover
title_p = find('PROJECT TITLE')
for r in title_p.runs:
    if '[Enter the approved project title]' in r.text:
        r.text = r.text.replace('[Enter the approved project title]', C.TITLE)
        break
else:
    sys.exit('title placeholder not found in one run')
for prefix, text in C.COVER.items():
    set_text(find(prefix), text)

# the page break after the cover becomes part of its last line, so a longer cover cannot push an
# empty break paragraph onto a page of its own
date_p = find('Date of Submission:')
break_p = Paragraph(date_p._p.getnext(), doc._body)
if 'w:type="page"' not in break_p._p.xml or break_p.text.strip():
    sys.exit('expected the cover page-break paragraph after the date line')
delete(break_p)
date_p.runs[-1].add_break(WD_BREAK.PAGE)

# 2. instruction page: from its heading up to and including the page break before the abstract
start = find_heading('Document Preparation Instructions')
abstract_h = find_heading('Abstract')
elm = start._p
while elm is not abstract_h._p:
    nxt = elm.getnext()
    elm.getparent().remove(elm)
    elm = nxt

# 3. abstract, keywords, table of contents
replace('[Write 200–300 words', [('p', t) for t in C.ABSTRACT])
replace('[Provide 4–6 keywords', [('p', C.KEYWORDS)])
toc_h = find_heading('Table of Contents')
page_break_before(toc_h)
toc_instr = find('In Microsoft Word, use References')
elm = toc_instr._p.getnext()
while Paragraph(elm, doc._body).style.name == 'List Bullet':
    nxt = elm.getnext()
    elm.getparent().remove(elm)
    elm = nxt
fld = new_after(toc_instr._p)
r = fld.add_run()
for tag, attr in (('w:fldChar', 'begin'), ('w:instrText', None), ('w:fldChar', 'separate')):
    e = OxmlElement(tag)
    if attr:
        e.set(qn('w:fldCharType'), attr)
    else:
        e.set(qn('xml:space'), 'preserve')
        e.text = 'TOC \\o "1-2" \\h \\z \\u'
    r._r.append(e)
fld.add_run('Right-click and choose Update Field to build the table of contents.')
r3 = fld.add_run()
e = OxmlElement('w:fldChar')
e.set(qn('w:fldCharType'), 'end')
r3._r.append(e)
delete(toc_instr)
page_break_before(find_heading('1. Introduction'))

# 4. chapters
replace('[Introduce the application domain', C.INTRO_BACKGROUND)
replace('[State the specific problem', C.INTRO_PROBLEM)
replace('[Explain the practical need', C.INTRO_MOTIVATION)
replace('[Define what is included', C.INTRO_SCOPE)
replace('[Briefly describe the content of each chapter', C.INTRO_ORGANISATION)
replace('Review 15 relevant research papers', C.LIT_INTRO)
replace('[Compare the papers by approach', C.LIT_SYNTHESIS)
replace('Each student must prepare an independent analysis', C.GAP_INTRO)
replace('[Summarise the main gap', C.GAP_CONSOLIDATED)
replace('Write 4–6 specific, measurable', C.OBJ_INTRO)
for i, text in enumerate(C.OBJECTIVES, 1):
    p = find(f'Objective {i}: [')
    p.style = doc.styles['Normal']
    set_text(p, text, bold_lead=True)
replace('[List expected deliverables', C.OBJ_OUTCOMES)
replace('Limit this section to approximately one page', C.NOVELTY_INTRO)
replace('[State the central novelty claim', C.NOVELTY_CLAIM)
replace('[Describe the users, inputs', C.SYS_OVERVIEW)
replace('[Insert the AWS architecture diagram here', [('fig', 'aws')])
delete(find('Figure 1. Proposed AWS Cloud Architecture.'))
replace('[Explain the role of each component', C.SYS_AWS)
replace('[Insert the complete project/system workflow diagram', [('fig', 'system')])
delete(find('Figure 2. Complete System Architecture.'))
replace('[Explain how the complete workflow operates', C.SYS_COMPLETE)
replace('[Add a numbered workflow', C.SYS_WORKFLOW)
replace('[Explain IAM least privilege', C.SYS_SECURITY)
replace('Complete this section if the project uses a dataset', C.DATA_INTRO)
replace('[Describe train/validation/test split', C.DATA_PREP)
replace('List only services relevant to this project', C.AWS_INTRO)
replace('[Describe region, VPC/subnets', C.AWS_NETWORK)
replace('[Provide an estimated monthly cost', C.AWS_COST)
replace('[Describe the phases: requirements', C.METHOD + [('fig', 'gnn')])
for i, text in enumerate(C.STEPS, 1):
    p = find(f'Step {i} –')
    p.style = doc.styles['Normal']
    set_text(p, text, bold_lead=True)
replace('[Specify relevant measures', C.METRICS)
replace('For Phase I, document completed work', C.IMPL_INTRO)
replace('[Insert relevant console, application', C.SCREENSHOTS)
replace('[Repository name should follow', C.REPO_TEXT)
struct = find('Suggested structure:')
set_text(struct, 'The main folders are:')
elm = struct._p.getnext()
while Paragraph(elm, doc._body).style.name == 'List Bullet':
    nxt = elm.getnext()
    elm.getparent().remove(elm)
    elm = nxt
render([('b', C.REPO_STRUCTURE)], struct._p)
replace('[Insert actual results', C.RESULTS)
replace('[Compare the proposed system with a justified baseline', C.BASELINE)
replace('[Summarise which objectives have been achieved', C.RESULTS_SUMMARY)
replace('[Interpret the findings', C.DISCUSSION)
replace('[Document technical, dataset', C.LIMITATIONS)
replace('For team projects, clearly document', C.TEAM_INTRO)
replace('[Document main, develop, and feature/studentX', C.WORKFLOW)
replace("[Add each student's GitHub profile", C.EVIDENCE)
replace('[Summarise the problem, proposed cloud solution', C.CONCLUSION)
replace('[List realistic next steps', C.FUTURE)

# 5. references
delete(find('Use APA style consistently'))
for i in range(1, 16):
    set_text(find(f'[{i}] [Author(s)'), f'[{i}] {C.REFERENCES[i - 1]}')
last = find('[15] ')._p
for i in range(16, len(C.REFERENCES) + 1):
    para = new_after(last)
    src_ppr = last.find(qn('w:pPr'))
    if src_ppr is not None:
        para._p.insert(0, copy.deepcopy(src_ppr))
    add_text(para, f'[{i}] {C.REFERENCES[i - 1]}')
    last = para._p
delete(find('[Add additional references'))

# 6. tables, in template order
T = doc.tables
fill_table(T[0], C.LIT_ROWS)
fill_table(T[1], C.GAP_ROWS)
fill_table(T[2], C.NOVELTY_ROWS)
fill_table(T[3], C.DATA_ROWS)
fill_table(T[4], C.AWS_ROWS)
fill_table(T[5], C.TOOLS_ROWS)
fill_table(T[6], C.MODULE_ROWS)
fill_table(T[7], C.TEST_ROWS)
fill_table(T[8], C.RISK_ROWS)
fill_table(T[9], C.TEAM_ROWS, header=C.TEAM_HEADER)
fill_table(T[10], C.TIMELINE_ROWS, header=C.TIMELINE_HEADER)

# 7. appendix B checklist
for p in paragraphs():
    if p.text.startswith('☐'):
        for key, ok in C.CHECKLIST.items():
            if key in p.text:
                set_text(p, ('☑' + p.text[1:]) if ok is True else f'{p.text} (Not yet: {ok}.)')
                break
        else:
            sys.exit(f'checklist line without a decision: {p.text}')

# 8. nothing of the template's guidance may survive
left = [p.text for p in paragraphs() if re.search(r'\[(?!TO CONFIRM)[A-Z][^\]]*\]', p.text) and not re.match(r'^\[\d+\] ', p.text)]
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            if re.search(r'\[(?!TO CONFIRM)[A-Z][a-z][^\]]*\]', cell.text):
                left.append(cell.text)
if left:
    sys.exit('template placeholders left:\n  ' + '\n  '.join(left[:20]))

def section_break_after(elm, landscape: bool):
    """Ends the current section at a new paragraph placed after elm; the section is landscape or portrait."""
    p = OxmlElement('w:p')
    elm.addnext(p)
    ppr = OxmlElement('w:pPr')
    p.append(ppr)
    sect = copy.deepcopy(BODY.find(qn('w:sectPr')))
    pg = sect.find(qn('w:pgSz'))
    w, h = int(pg.get(qn('w:w'))), int(pg.get(qn('w:h')))
    if landscape:
        pg.set(qn('w:w'), str(max(w, h)))
        pg.set(qn('w:h'), str(min(w, h)))
        pg.set(qn('w:orient'), 'landscape')
    ppr.append(sect)
    return p


for heading, key in (('5.2 Diagram 1 – AWS Cloud Architecture (Mandatory)', 'aws'), ('5.3 Diagram 2 – Complete System Architecture (Mandatory)', 'system')):
    h = find_heading(heading)
    section_break_after(h._p.getprevious(), landscape=False)      # previous text stays portrait
    caption = h._p.getnext().getnext()                               # image paragraph, then caption
    if 'Figure' not in Paragraph(caption, doc._body).text:
        sys.exit(f'caption not found after {heading}')
    section_break_after(caption, landscape=True)                     # heading + figure are landscape


doc.core_properties.title = C.TITLE
doc.core_properties.subject = 'BCSE355L Cloud Architecture Design – AWS Cloud Project Report'
doc.core_properties.author = 'Naveen Kumar, Kabilesh SN'
doc.save(OUT)
print(f'wrote {OUT} with {len(fig_numbers)} figures')
