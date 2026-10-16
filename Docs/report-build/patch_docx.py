"""Replaces exact text in the finished report in place, keeping every other edit made in Word.

    python patch_docx.py "old text" "new text" ["old text" "new text" ...]

Matches inside a single run are replaced in that run (formatting kept). A match split across runs
is rewritten into the first run of the span. Every pair must match at least once, or nothing is saved.
"""
import sys
from pathlib import Path

import docx

DOC = Path(__file__).resolve().parents[1] / 'AWS_Cloud_Project_Report_Final.docx'
pairs = list(zip(sys.argv[1::2], sys.argv[2::2]))
d = docx.Document(DOC)


def paragraphs():
    yield from d.paragraphs
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for s in d.sections:
        for part in (s.header, s.footer):
            yield from part.paragraphs


counts = {old: 0 for old, _ in pairs}
seen = set()
for p in paragraphs():
    if p._p in seen:
        continue
    seen.add(p._p)
    for old, new in pairs:
        while old in p.text:
            runs = p.runs
            done = False
            for r in runs:
                if old in r.text:
                    r.text = r.text.replace(old, new, 1)
                    done = True
                    break
            if not done:
                # spans runs: locate the span and collapse it into its first run
                text, starts = '', []
                for r in runs:
                    starts.append(len(text))
                    text += r.text
                i = text.index(old)
                j = i + len(old)
                first = max(k for k, s in enumerate(starts) if s <= i)
                last = max(k for k, s in enumerate(starts) if s < j)
                head = runs[first].text[: i - starts[first]]
                tail = runs[last].text[j - starts[last]:]
                runs[first].text = head + new + (tail if last != first else '')
                for r in runs[first + 1: last]:
                    r.text = ''
                if last != first:
                    runs[last].text = ''
                    runs[first].text = head + new + tail
            counts[old] += 1

missing = [o for o, c in counts.items() if c == 0]
if missing:
    sys.exit(f'not found, nothing saved: {missing}')
d.save(DOC)
for o, c in counts.items():
    print(f'{c} x  {o!r}')
